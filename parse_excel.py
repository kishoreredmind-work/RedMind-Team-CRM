import openpyxl
import json
import uuid

wb = openpyxl.load_workbook("d:/Redmin Team CRM/RedMind project tracker.xlsx")
projects = []

# 1. Website Projects
sheet = wb['Website Projects']
rows = list(sheet.iter_rows(values_only=True))
headers = rows[0]
for row in rows[1:]:
    if not row or not any(row):
        continue
    data = dict(zip(headers, row))
    proj_name = data.get('Client / Project')
    if not proj_name:
        continue
    
    status_str = str(data.get('Status', '')).strip().lower()
    status = 'Completed' if 'complete' in status_str else 'In Progress' if 'ongoing' in status_str else 'New'
    
    projects.append({
        'name': str(proj_name).strip(),
        'description': f"Type: {data.get('Type') or 'Website'}. Notes: {data.get('Notes / Additional Links') or ''}",
        'status': status,
        'website': data.get('Website') or '',
        'category': 'Website'
    })

# 2. Graphic Design Projects
sheet = wb['Graphic Design Projects']
rows = list(sheet.iter_rows(values_only=True))
headers = rows[0]
for row in rows[1:]:
    if not row or not any(row):
        continue
    data = dict(zip(headers, row))
    proj_name = data.get('Client / Project')
    if not proj_name:
        continue
    
    status_str = str(data.get('Status', '')).strip().lower()
    status = 'Completed' if 'complete' in status_str else 'In Progress' if 'ongoing' in status_str else 'New'
    
    projects.append({
        'name': str(proj_name).strip(),
        'description': f"Service(s): {data.get('Service(s)') or 'Graphic Design'}. Notes: {data.get('Notes') or ''}",
        'status': status,
        'website': '',
        'category': 'Graphic Design'
    })

# 3. AI projects
sheet = wb['AI projects']
rows = list(sheet.iter_rows(values_only=True))
headers = rows[0]
for row in rows[1:]:
    if not row or not any(row):
        continue
    data = dict(zip(headers, row))
    proj_name = data.get('Client / Project')
    if not proj_name:
        continue
    
    status_str = str(data.get('Status', '')).strip().lower()
    status = 'Completed' if 'complete' in status_str else 'In Progress' if 'ongoing' in status_str else 'New'
    
    projects.append({
        'name': str(proj_name).strip(),
        'description': f"AI Project. Type: {data.get('Type') or ''}. Service(s): {data.get('Service(s)') or ''}. POC: {data.get('POC') or ''}",
        'status': status,
        'website': '',
        'category': 'AI'
    })

with open("d:/Redmin Team CRM/projects_to_migrate.json", "w", encoding="utf-8") as f:
    json.dump(projects, f, indent=4)

print(f"Parsed {len(projects)} projects and written to projects_to_migrate.json")
